#!/usr/bin/env python3
"""Build registered multilingual article pages from the canonical workbook.

Article prose and publication metadata come only from vizier.xlsx.
Only marked front-page/list blocks and the registered article are written.
Existing text, articles, menus and banners are otherwise preserved.
"""
from pathlib import Path
from html import escape
from datetime import datetime
import re, json
import openpyxl
from bs4 import BeautifulSoup

REPO = Path(__file__).resolve().parent.parent
wb = openpyxl.load_workbook(REPO / "nl/_data/vizier.xlsx", data_only=True)
if "Artikelinhoud" not in wb.sheetnames:
    raise SystemExit(0)
matrix = list(wb["Knopen"].values)
headers = matrix[1]
rows = [{h: row[i] for i, h in enumerate(headers) if h} for row in matrix[2:]]
articles = {}
for row in wb["Artikelinhoud"].iter_rows(min_row=3, values_only=True):
    if row[1]:
        articles.setdefault(row[1], []).append((int(row[2]), row[3], row[4]))
template = (REPO / "scripts/templates/xlsx-artikel.html").read_text()
settings = {}
if "Artikeltaal" in wb.sheetnames:
    for lang, slug, raw in wb["Artikeltaal"].iter_rows(min_row=3, max_col=3, values_only=True):
        if slug:
            settings[slug] = (lang, json.loads(raw))
LANGUAGE_NAMES={"nl":"Nederlands","en":"English","de":"Deutsch","fr":"Français",
                "es":"Español","it":"Italiano","pt":"Português","ru":"Русский"}
published={}
for slug in articles:
    matching=[r for r in rows if str(r.get("URL","")).endswith("/"+slug+".html")
              and r.get("Status")=="live" and r.get("Actief") is not False]
    if matching:
        published[slug]=matching[0]

def replace_marked(path, key, markup, anchor):
    text = path.read_text()
    start, end = f"<!-- XLSX {key} START -->", f"<!-- XLSX {key} END -->"
    block = f"{start}\n{markup}\n{end}"
    if start in text:
        text = re.sub(re.escape(start) + r".*?" + re.escape(end), lambda _: block, text, flags=re.S)
    else:
        assert anchor in text, f"Missing insertion anchor in {path}"
        text = text.replace(anchor, anchor + "\n" + block, 1)
    path.write_text(text)

for slug, blocks in articles.items():
    assert re.fullmatch(r"[a-z0-9-]+", slug), "Unsafe article slug"
    matching = [r for r in rows if str(r.get("URL","")).endswith("/"+slug+".html")]
    assert matching, f"No matrix row for {slug}"
    record = matching[0]
    if record.get("Status") != "live" or record.get("Actief") is False:
        continue
    lang,ui=settings[slug]
    assert lang==record["Taal"]
    route=record["URL"]
    assert re.fullmatch(r"[a-z0-9/-]+\.html",route)
    family=[r for r in published.values() if r.get("Extra 1")==record.get("Extra 1")]
    family.sort(key=lambda r:list(LANGUAGE_NAMES).index(r["Taal"]))
    alternates="\n".join(f'<link rel="alternate" hreflang="{r["Taal"]}" href="https://openvizier.org/{r["Taal"]}/{r["URL"]}">' for r in family)
    language_links=" ".join(f'<a href="/{r["Taal"]}/{r["URL"]}" lang="{r["Taal"]}" hreflang="{r["Taal"]}"'
        +(' aria-current="page"' if r["Taal"]==lang else '')
        +f'>{LANGUAGE_NAMES[r["Taal"]]}</a>' for r in family)
    blocks.sort()
    first = {slot: html for _, slot, html in blocks if slot != "body"}
    raw_body = "\n".join(html for _, slot, html in blocks if slot == "body")
    body = BeautifulSoup(raw_body, "html.parser")
    toc = []
    for heading in body.find_all(["h1", "h2"]):
        if not heading.get_text(strip=True):
            heading.decompose()
            continue
        if heading.name == "h1":
            heading.name = "h2"
            toc.append(f'<a href="#{escape(heading["id"])}">{escape(heading.get_text(" ", strip=True))}</a>')
        else:
            heading.name = "h2" if heading.get("id") == "leeswijzer" else "h3"
    for table in body.find_all("table"):
        wrap = body.new_tag("div", attrs={"class": "table-scroll", "tabindex": "0", "role": "region",
            "aria-label": ui["table_aria"]})
        table.wrap(wrap)
    in_sources = False
    for el in list(body.children):
        if not getattr(el, "name", None):
            continue
        if el.get("id") == "bronnen":
            in_sources = True
        if in_sources and el.name == "p":
            el["class"] = "source-entry"
            for text in list(el.find_all(string=True)):
                if text.parent.name == "a":
                    continue
                chunks = re.split(r"(https?://[^\s<>]+)", str(text))
                if len(chunks) < 2:
                    continue
                for chunk in chunks:
                    if chunk.startswith(("http://", "https://")):
                        a = body.new_tag("a", href=chunk, target="_blank", rel="noopener noreferrer")
                        a.string = chunk
                        text.insert_before(a)
                    else:
                        text.insert_before(chunk)
                text.extract()
    assert re.sub(r"\s+", "", BeautifulSoup(raw_body, "html.parser").get_text()) == re.sub(r"\s+", "", body.get_text())
    title = record["Naam"]
    subtitle = record.get("Ondertitel") or ""
    hero = str(record["Hero"]).replace("../", "/", 1)
    url = f"https://openvizier.org/{lang}/{route}"
    values = {
        "TITLE": escape(title), "DESCRIPTION": escape(subtitle, quote=True), "CANONICAL": url,
        "HERO": hero, "HERO_ABSOLUTE": "https://openvizier.org" + hero,
        "SUBTITLE": first["subtitle"], "DATE": first["date"], "BODY": str(body),
        "TOC": "\n".join(toc), "PUBLICATION_DATE": str(record["Datum publicatie"]),
        "AUTHOR": escape(record.get("Auteur") or ""),
        "LANG":lang, "ALTERNATES":alternates, "LANGUAGE_LINKS":language_links,
        **{"UI_"+k.upper():escape(v,quote=True) for k,v in ui.items()},
    }
    page = template
    for key, value in values.items():
        page = page.replace("{{" + key + "}}", value)
    assert "{{" not in page
    output=REPO/lang/route
    output.parent.mkdir(parents=True,exist_ok=True)
    output.write_text(page)
    label=ui["date"]
    banner = f'''<a class="dgk-top" href="/{lang}/{route}" style="background:#1a3347;border-left-color:#bf9b56;">
  <div class="dgk-top__grid">
    <div class="dgk-top__img"><img src="{hero}" alt="{escape(title)}" loading="eager" decoding="async"></div>
    <div class="dgk-top__txt">
      <p class="dgk-top__tag">{escape(ui["new"])} · {escape(ui["forecast"].upper())} · {escape(label.upper())}</p>
      <h2 class="dgk-top__h">{escape(title)}</h2>
      <p class="dgk-top__l">{escape(subtitle)}</p>
      <span class="dgk-top__k">{escape(ui["read"])} →</span>
    </div>
  </div>
</a>'''
    replace_marked(REPO / lang / "index.html", slug, banner,
        '<script src="/assets/menu-loader.js" defer></script>')
    listing = f'''<article class="wo-item">
  <div class="wo-item__layout">
    <div class="wo-item__visual"><a href="{slug}.html"><img src="{hero}" alt="{escape(title)}" loading="eager"></a></div>
    <div class="wo-item__tekst">
      <p class="wo-item__datum">{escape(label)} · {escape(ui["forecast"])}</p>
      <h2 class="wo-item__titel"><a href="{slug}.html">{escape(title)}</a></h2>
      <p class="wo-item__lead">{escape(subtitle)}</p>
    </div>
  </div>
</article>'''
    listing_path = output.parent / "index.html"
    match = re.search(r'<(?:section|div)\b[^>]*class=["\'][^"\']*\bwo-lijst\b[^"\']*["\'][^>]*>', listing_path.read_text())
    assert match, "Missing chronological article list"
    replace_marked(listing_path, slug, listing, match.group())
    print(f"Built {slug}: {len(blocks)} content blocks; front page and article list updated.")
